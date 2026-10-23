"""
Lee la estructura de capas editada a mano en el Excel (hoja «Elementos») y la guarda en datos/estructura.json,
que 03b_aplicar_estructura.py aplica a las capas base en cada actualización.

Cada fila de nivel 1-3 del Excel es una capa, subcapa o sub-subcapa (con su código, título, tipo y descripción),
y cada fila de nivel 4, un elemento. Los elementos se reconocen por nombre y coordenadas en las capas base
(datos/build/base/, que escribe 03_construir_capas.py) y se guardan por su identificador. Si una subcapa
recibe todos los elementos de una capa base, se guarda como «todos», de modo que los elementos nuevos de
esa capa base entran solos en cada actualización. Las filas que no están en las capas base son elementos
añadidos a mano y se guardan con sus datos.

Solo se ejecuta cuando se edita la estructura en el Excel:
    python estructura_desde_excel.py <excel>
"""

import collections, json, os, re, sys, unicodedata
from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
BASE = os.path.join(RAIZ, "datos", "build", "base")
SALIDA = os.path.join(RAIZ, "datos", "estructura.json")

# Decisiones del usuario sobre la edición del 23-10-2026
DESCARTADAS = {"otros_restaurantes", "carriles_bici"}   # quedaron bajo Delicatessen y Miradores al borrar su cabecera
TITULOS = {"Wllnes & spa": "Wellness & spa",
           "Experiencias (quitar las que no son de marbella)": "Experiencias"}
SOLO_TERMINO = {"Experiencias"}                          # sin los elementos fuera del término municipal
PENDIENTE = "Pendiente de completar"
COLORES_NUEVOS = ["#00838F", "#558B2F", "#6A1B9A", "#EF6C00", "#5E35B1", "#00695C"]


def slug(t):
    t = unicodedata.normalize("NFD", t).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")


def limpio(t):
    t = " ".join(str(t).split())
    return TITULOS.get(t, t)


cat = json.load(open(os.path.join(BASE, "catalogo.json"), encoding="utf8"))
base_sub, base_grp, por_clave, ids_de = {}, {}, collections.defaultdict(list), {}
for g in cat["grupos"]:
    for s in g["subcapas"]:
        base_sub[s["id"]] = s
        base_grp[s["id"]] = g
        fs = json.load(open(os.path.join(BASE, s["id"] + ".geojson"), encoding="utf8"))["features"]
        ids_de[s["id"]] = {f["id"] for f in fs}
        for f in fs:
            p = f["properties"]
            por_clave[(p["nombre"], p.get("coord") or "Sin ubicación")].append((s["id"], f["id"]))

# -- leer el árbol del Excel ------------------------------------------------------------------------------
ws = load_workbook(sys.argv[1])["Elementos"]
capas, pila = [], {}
for r in ws.iter_rows(min_row=5, values_only=True):
    cod, nivel = r[0], r[6]
    if nivel == 1:
        capa = {"titulo": limpio(r[1]), "descripcion": r[12], "web": r[15], "subcapas": []}
        capas.append(capa)
    elif nivel in (2, 3):
        titulo = limpio(r[2] if nivel == 2 else r[3])
        nodo = {"titulo": titulo, "nivel": nivel, "tipo": r[7], "descripcion": r[12], "filas": [],
                "padre": pila[2]["titulo"] if nivel == 3 else None}
        pila[nivel] = nodo
        capa["subcapas"].append(nodo)
    elif nivel == 4:
        nodo["filas"].append(r)
    elif r[1] == "Total":
        break

# -- reconocer los elementos ----------------------------------------------------------------------------
usados = set()
for capa in capas:
    ids_tit = {}
    for s in capa["subcapas"]:
        cands = [[c for c in por_clave.get((r[4], r[5]), []) if c[0] not in DESCARTADAS] for r in s["filas"]]
        frec = collections.Counter(sid for cs in cands for sid in {c[0] for c in cs})
        def prioridad(sid):
            return (base_sub[sid]["titulo"] == s["titulo"], frec[sid] == len(ids_de[sid]), frec[sid])
        de, nuevos = collections.defaultdict(set), []
        for r, cs in zip(s["filas"], cands):
            if cs:
                sid, fid = max(cs, key=lambda c: prioridad(c[0]))
                de[sid].add(fid)
            elif not por_clave.get((r[4], r[5])):          # no está en ninguna capa base: añadido a mano
                lat_lon = None if r[5] in (None, "Sin ubicación") else [float(x) for x in r[5].split(",")]
                nuevos.append({"nombre": r[4], "coord": lat_lon, "direccion": r[12], "fuente": r[13] or "AMMA Consulting",
                               "aviso": r[14], "ficha": r[15]})
        principal = max(de, key=prioridad) if de else None
        s["de"] = [{"capa": sid, "todos": True} if ids == ids_de[sid] else {"capa": sid, "ids": sorted(ids)}
                   for sid, ids in sorted(de.items(), key=lambda x: -len(x[1]))]
        s["nuevos"] = nuevos
        # identificador: el de la capa base si la reproduce, para conservar enlaces e iconos; si no, el título
        # (precedido del de la subcapa madre en las sub-subcapas: «hospitales_privados»)
        reproduce = principal and principal not in usados and (
            base_sub[principal]["titulo"] == s["titulo"] or (len(de) == 1 and s["de"][0].get("todos")))
        sid = principal if reproduce else slug(s["titulo"])
        if not reproduce and s["padre"]:
            sid = ids_tit[s["padre"]] + "_" + sid
        while sid in usados:
            sid += "_2"
        usados.add(sid)
        s["id"] = ids_tit[s["titulo"]] = sid
        s["icono"] = principal
        s["complementaria"] = s.pop("tipo") == "Complementaria"
        if s["titulo"] in SOLO_TERMINO:
            s["solo_termino"] = True
        if not s["filas"] and not any(x["padre"] == s["titulo"] for x in capa["subcapas"]):
            s["pendiente"] = PENDIENTE
        del s["filas"]

# padres por id y datos de capa (color e identificador de la capa base de la que procede la mayoría)
colores = iter(COLORES_NUEVOS)
ids_capa = set()
for capa in capas:
    tit = {s["titulo"]: s["id"] for s in capa["subcapas"]}
    for s in capa["subcapas"]:
        s["padre"] = tit[s["padre"]] if s["padre"] else None
        s["nivel"] = 3 if s["padre"] else 2
    origen = collections.Counter(base_grp[s["icono"]]["id"] for s in capa["subcapas"] if s["icono"])
    gb = next((g for g in cat["grupos"] if g["titulo"] == capa["titulo"]), None)
    if not gb and origen:
        gid = origen.most_common(1)[0][0]
        gb = next(g for g in cat["grupos"] if g["id"] == gid)
        gb = gb if gid not in ids_capa and not any(c["titulo"] == gb["titulo"] for c in capas) else None
    capa["id"] = gb["id"] if gb and gb["id"] not in ids_capa else slug(capa["titulo"])
    ids_capa.add(capa["id"])
    capa["color"] = gb["color"] if gb else next(colores)

# Elementos quitados a mano: los de las capas base que no están en ninguna subcapa (para que, en las
# actualizaciones, solo se avise de los elementos nuevos que no tienen sitio en el árbol)
asignados = collections.defaultdict(set)
todas = set()
for capa in capas:
    for s in capa["subcapas"]:
        for d in s["de"]:
            todas.add(d["capa"]) if d.get("todos") else asignados[d["capa"]].update(d["ids"])
excluidos = {sid: sorted(ids_de[sid] - ids) for sid, ids in asignados.items() if sid not in todas and ids_de[sid] - ids}
descartadas = sorted(sid for sid in ids_de if sid not in todas and sid not in asignados)

json.dump({"fuente": os.path.basename(sys.argv[1]), "capas": capas, "excluidos": excluidos,
           "descartadas": descartadas}, open(SALIDA, "w", encoding="utf8"), ensure_ascii=False, indent=1)
print("descartadas:", descartadas, "· excluidos:", {k: len(v) for k, v in excluidos.items()})
print(SALIDA)
for c in capas:
    print(c["id"], c["titulo"], c["color"])
    for s in c["subcapas"]:
        print("   " + "   " * (s["nivel"] - 2), s["id"], "·", s["titulo"],
              [d["capa"] + ("" if d.get("todos") else "(%d)" % len(d["ids"])) for d in s["de"]],
              "+%d nuevos" % len(s["nuevos"]) if s["nuevos"] else "", s.get("pendiente") or "")
