"""
Codifica la estructura del visor y la publica en Excel.

1. Da a cada elemento de web/data/*.geojson su código (capa.subcapa[.sub-subcapa].elemento, p. ej. 01.02.005,
   numerado por orden alfabético dentro de su subcapa) y una sola coordenada «lat, lon» (WGS84):
   el punto; en polígonos, el marcador de la web o su centro (siempre dentro del polígono); en líneas,
   el punto a mitad de recorrido. Los dos quedan en las propiedades `codigo` y `coord`, que muestra el visor.
2. Escribe web/estructura-capas.xlsx (se descarga desde el visor) con el árbol capa › subcapa › sub-subcapa
   › elemento, plegable por niveles, y una hoja de resumen por capa. Fuera de GitHub Actions deja además
   una copia en la carpeta del proyecto con el nombre del encargo.

Uso: python 05_estructura_excel.py (después de 03_construir_capas.py)
"""

import json, os, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from shapely.geometry import shape
from shapely.ops import linemerge

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DATA = os.path.join(RAIZ, "web", "data")
SALIDA = os.path.join(RAIZ, "web", "estructura-capas.xlsx")
COPIA = os.path.join(RAIZ, "Visor de Turismo SITMA - Estructura de capas y subcapas.xlsx")
d = json.load(open(os.path.join(DATA, "catalogo.json"), encoding="utf8"))

TURQ = "37A0B4"
F = "Arial"
borde = Border(bottom=Side(style="thin", color="D0D7DE"))

cab = ["Código", "Capa", "Subcapa", "Sub-subcapa", "Elemento", "Coordenadas (lat, lon · WGS84)", "Nivel", "Tipo",
       "Geometría", "Elementos en el mapa", "Sin ubicación", "Recursos en la web", "Descripción / dirección",
       "Fuente", "Pendiente / aviso", "Página de la web"]
C = {k: i for i, k in enumerate(["cod", "capa", "sub", "subsub", "elem", "coord", "nivel", "tipo", "geom",
                                 "n", "sin", "nweb", "desc", "fuente", "pend", "web"], 1)}
L = {k: get_column_letter(v) for k, v in C.items()}
CENTRO = {C["cod"], C["nivel"], C["geom"], C["n"], C["sin"], C["nweb"]}
WRAP = {C["desc"], C["fuente"], C["pend"], C["web"]}
GEOM = {"Point": "punto", "Polygon": "polígono", "MultiPolygon": "polígono", "LineString": "línea",
        "MultiLineString": "línea"}


def tint(hexc, f):
    r, g, b = (int(hexc[i:i + 2], 16) for i in (1, 3, 5))
    return "%02X%02X%02X" % tuple(int(c + (255 - c) * f) for c in (r, g, b))


def ll(p):
    return "%.6f, %.6f" % (p[1], p[0])


def coordenadas(ft):
    g, p = ft["geometry"], ft["properties"]
    if not g:
        return None
    if g["type"] == "Point":
        return ll(g["coordinates"])
    if "pin" in p:
        return ll(p["pin"])
    geo = shape(g)
    if g["type"] in ("LineString", "MultiLineString"):
        if g["type"] == "MultiLineString":
            geo = linemerge(geo)
            if geo.geom_type == "MultiLineString":
                geo = max(geo.geoms, key=lambda x: x.length)
        pt = geo.interpolate(0.5, normalized=True)
    else:
        pt = geo.centroid
        if not geo.contains(pt):
            pt = geo.representative_point()
    return ll([pt.x, pt.y])


wb = Workbook()
ws = wb.active
ws.title = "Estructura"
ws.sheet_properties.outlinePr.summaryBelow = False

ws["A1"] = "Estructura de capas y subcapas del visor"
ws["A1"].font = Font(name=F, size=16, bold=True, color=TURQ)
ws["A2"] = ("Marbella · AMMA Consulting · catálogo generado el %s. Use los botones 1-2-3-4 del margen "
            "izquierdo para plegar o desplegar las ramas (4 = todo desplegado, con cada elemento y sus coordenadas)."
            % d["generado"])
ws["A2"].font = Font(name=F, size=10, italic=True, color="555555")

HR = 4
for i, c in enumerate(cab, 1):
    x = ws.cell(HR, i, c)
    x.font = Font(name=F, bold=True, color="FFFFFF")
    x.fill = PatternFill("solid", fgColor=TURQ)
    x.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center" if i in CENTRO else "left")
ws.row_dimensions[HR].height = 32

r = HR + 1
filas_capa = []
n_elem = 0
for gi, g in enumerate(d["grupos"], 1):
    hijos = {}
    for s in g["subcapas"]:
        hijos.setdefault(s["padre"], []).append(s)
    rc = r
    filas_capa.append(rc)
    vals = {"cod": "%02d" % gi, "capa": g["titulo"], "nivel": 1, "tipo": "Capa", "desc": g["descripcion"],
            "web": g["web"]}
    for k, i in C.items():
        x = ws.cell(rc, i, vals.get(k))
        x.font = Font(name=F, bold=True, color="FFFFFF")
        x.fill = PatternFill("solid", fgColor=g["color"].lstrip("#"))
        x.alignment = Alignment(vertical="top", wrap_text=(i == C["desc"]),
                                horizontal="center" if i in CENTRO else "left")
    r += 1
    first = r

    def escribe(s, nivel, cod):
        global r, n_elem
        s["codigo"] = cod
        vals = {"cod": cod, "nivel": nivel, "tipo": "Complementaria" if s["complementaria"] else "De la web",
                "geom": s["geometria"], "n": s["n"], "sin": s["sin_ubicacion"], "nweb": s["n_web"],
                "desc": s["descripcion"], "fuente": s["fuente"], "pend": s["pendiente"],
                "web": "\n".join(s["web"]) if s["web"] else None,
                ("sub" if nivel == 2 else "subsub"): s["titulo"]}
        fill = PatternFill("solid", fgColor=tint(g["color"], 0.85 if nivel == 2 else 0.93))
        for k, i in C.items():
            x = ws.cell(r, i, vals.get(k))
            x.font = Font(name=F, bold=(nivel == 2 and k == "sub"),
                          italic=s["complementaria"] and k in ("sub", "subsub", "tipo"),
                          color="C0392B" if k == "pend" else "222222")
            x.fill = fill
            x.border = borde
            x.alignment = Alignment(vertical="top", wrap_text=i in WRAP,
                                    horizontal="center" if i in CENTRO else "left")
        ws.row_dimensions[r].outlineLevel = nivel - 1
        r += 1
        # Elementos de la subcapa: código y coordenada, en el Excel y en el GeoJSON del visor
        ruta = os.path.join(DATA, s["id"] + ".geojson")
        fc = json.load(open(ruta, encoding="utf8"))
        orden = sorted(fc["features"], key=lambda f: f["properties"]["nombre"].lower())
        for k, ft in enumerate(orden, 1):
            p = ft["properties"]
            gm = ft["geometry"]
            p["codigo"] = "%s.%03d" % (cod, k)
            p.pop("coord", None)
            if gm:
                p["coord"] = coordenadas(ft)
            vals = {"cod": p["codigo"], "elem": p["nombre"], "coord": p.get("coord") or "Sin ubicación",
                    "nivel": 4, "tipo": "Elemento", "geom": GEOM[gm["type"]] if gm else None,
                    "desc": p.get("direccion"), "fuente": p.get("fuente"), "pend": p.get("aviso"),
                    "web": p.get("ficha") or p.get("web")}
            for kk, i in C.items():
                x = ws.cell(r, i, vals.get(kk))
                x.font = Font(name=F, size=9, color="C0392B" if kk == "pend" else "444444")
                x.border = borde
                x.alignment = Alignment(vertical="top", horizontal="center" if i in CENTRO else "left")
            ws.row_dimensions[r].outlineLevel = nivel
            n_elem += 1
            r += 1
        json.dump(fc, open(ruta, "w", encoding="utf8"), ensure_ascii=False, separators=(",", ":"))
        for k, h in enumerate(hijos.get(s["id"], []), 1):
            escribe(h, nivel + 1, "%s.%02d" % (cod, k))

    g["codigo"] = "%02d" % gi
    for k, s in enumerate(hijos.get(None, []), 1):
        escribe(s, 2, "%02d.%02d" % (gi, k))
    last = r - 1
    # Totales de la capa: solo el nivel 2 (el nivel 3 son subconjuntos de una subcapa)
    for k in ("n", "sin", "nweb"):
        x = ws.cell(rc, C[k], "=SUMIFS(%s%d:%s%d,$%s%d:$%s%d,2)"
                    % (L[k], first, L[k], last, L["nivel"], first, L["nivel"], last))
        x.number_format = "#,##0"
    x = ws.cell(rc, C["sub"], '=COUNTIFS($%s%d:$%s%d,2)&" subcapas"' % (L["nivel"], first, L["nivel"], last))
    x.font = Font(name=F, italic=True, color="FFFFFF")

fin = r - 1
r += 1
ws.cell(r, C["capa"], "Total").font = Font(name=F, bold=True)
for k in ("n", "sin", "nweb"):
    x = ws.cell(r, C[k], "=SUMIFS(%s%d:%s%d,$%s%d:$%s%d,1)"
                % (L[k], HR + 1, L[k], fin, L["nivel"], HR + 1, L["nivel"], fin))
    x.font = Font(name=F, bold=True)
    x.number_format = "#,##0"
    x.alignment = Alignment(horizontal="center")
    x.border = Border(top=Side(style="medium", color=TURQ))
r += 2
notas = [
    "Notas",
    "Nivel 1 = capa; 2 = subcapa; 3 = subconjunto de una subcapa (p. ej. playas con Q de Calidad); 4 = elemento. "
    "Los totales de cada capa suman solo el nivel 2 para no contar dos veces.",
    "«De la web»: subcapa que reproduce una sección de turismo.marbella.es. «Complementaria» (en cursiva): "
    "subcapa añadida con otras fuentes (OSM, OpenRTA, REDIAM...).",
    "Coordenadas: una por ubicación, en grados decimales WGS84 (latitud, longitud). Puntos: las del visor. "
    "Polígonos: el marcador de la web si lo hay; si no, su centro (siempre dentro del polígono). Líneas: el punto "
    "a mitad de recorrido. Las ubicaciones aproximadas lo indican en «Pendiente / aviso». "
    "«Sin ubicación»: la web no la publica.",
    "Los códigos y las coordenadas son los mismos que muestra el visor al pinchar en cada elemento. "
    "Fuente: datos del visor (%s)." % d["generado"],
]
for i, t in enumerate(notas):
    x = ws.cell(r + i, 1, t)
    x.font = Font(name=F, bold=(i == 0), size=10 if i else 11, color="555555" if i else TURQ)

for i, a in enumerate([13, 20, 30, 28, 34, 30, 7, 14, 11, 12, 11, 12, 50, 40, 30, 50], 1):
    ws.column_dimensions[get_column_letter(i)].width = a
ws.freeze_panes = ws.cell(HR + 1, C["sub"])
ws.auto_filter.ref = "A%d:%s%d" % (HR, L["web"], fin)
ws.sheet_view.zoomScale = 90

# Hoja de resumen por capa
rs = wb.create_sheet("Resumen por capa")
rs["A1"] = "Resumen por capa"
rs["A1"].font = Font(name=F, size=16, bold=True, color=TURQ)
cab2 = ["Código", "Capa", "Subcapas (nivel 2)", "De la web", "Complementarias", "Sub-subcapas (nivel 3)",
        "Elementos en el mapa", "Sin ubicación", "Recursos en la web"]
for i, c in enumerate(cab2, 1):
    x = rs.cell(3, i, c)
    x.font = Font(name=F, bold=True, color="FFFFFF")
    x.fill = PatternFill("solid", fgColor=TURQ)
    x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
rs.row_dimensions[3].height = 32
limites = filas_capa + [fin + 1]
E = "Estructura!"
N, T = L["nivel"], L["tipo"]
for k, (a, b) in enumerate(zip(limites, limites[1:])):
    rr = 4 + k
    lo, hi = a + 1, b - 1
    rs.cell(rr, 1, "=%sA%d" % (E, a))
    rs.cell(rr, 2, "=%sB%d" % (E, a))
    rs.cell(rr, 3, "=COUNTIFS(%s$%s%d:$%s%d,2)" % (E, N, lo, N, hi))
    rs.cell(rr, 4, '=COUNTIFS(%s$%s%d:$%s%d,2,%s$%s%d:$%s%d,"De la web")' % (E, N, lo, N, hi, E, T, lo, T, hi))
    rs.cell(rr, 5, '=COUNTIFS(%s$%s%d:$%s%d,2,%s$%s%d:$%s%d,"Complementaria")'
            % (E, N, lo, N, hi, E, T, lo, T, hi))
    rs.cell(rr, 6, "=COUNTIFS(%s$%s%d:$%s%d,3)" % (E, N, lo, N, hi))
    for c, k2 in ((7, "n"), (8, "sin"), (9, "nweb")):
        rs.cell(rr, c, "=%s%s%d" % (E, L[k2], a))
    for c in range(1, 10):
        x = rs.cell(rr, c)
        x.font = Font(name=F)
        x.border = borde
        x.alignment = Alignment(horizontal="left" if c == 2 else "center")
        x.number_format = "#,##0"
tr = 4 + len(filas_capa)
rs.cell(tr, 2, "Total").font = Font(name=F, bold=True)
for c in range(3, 10):
    Lc = get_column_letter(c)
    x = rs.cell(tr, c, "=SUM(%s4:%s%d)" % (Lc, Lc, tr - 1))
    x.font = Font(name=F, bold=True)
    x.number_format = "#,##0"
    x.alignment = Alignment(horizontal="center")
    x.border = Border(top=Side(style="medium", color=TURQ))
for i, a in enumerate([9, 24, 13, 11, 15, 14, 13, 11, 12], 1):
    rs.column_dimensions[get_column_letter(i)].width = a
rs.freeze_panes = "A4"
wb.save(SALIDA)

# Los códigos de capa y subcapa también van al catálogo, para el árbol del visor
json.dump(d, open(os.path.join(DATA, "catalogo.json"), "w", encoding="utf8"), ensure_ascii=False, indent=1)

if not os.environ.get("GITHUB_ACTIONS"):
    try:
        shutil.copyfile(SALIDA, COPIA)
    except PermissionError:
        print("Aviso: la copia del proyecto está abierta en Excel; no se ha actualizado.")
print("%s: %d elementos" % (SALIDA, n_elem))
